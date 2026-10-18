import io
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

import streamlit as st
import fitz  # PyMuPDF
import pytesseract
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PptInches, Pt as PptPt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.units import inch

# Optional libraries. The app still starts if AI cleanup is not installed.
try:
    from langchain_core.prompts import PromptTemplate, FewShotPromptTemplate
    from langchain_groq import ChatGroq
    LANGCHAIN_AVAILABLE = True
except ImportError:
    LANGCHAIN_AVAILABLE = False

try:
    from pdf2docx import Converter as PDFConverter
    PDF2DOCX_AVAILABLE = True
except ImportError:
    PDF2DOCX_AVAILABLE = False


st.set_page_config(
    page_title="AveonRush Document Converter",
    page_icon="🔄",
    layout="wide",
)

st.title("🔄 AveonRush Document Converter")
st.write(
    "Convert files with format-specific tools. Native document conversion is used "
    "where possible; some cross-format conversions can only approximate the original layout."
)

CONVERSIONS = {
    "PDF → DOCX": (["pdf"], "docx"),
    "PDF → PPTX": (["pdf"], "pptx"),
    "PDF → TXT": (["pdf"], "txt"),
    "DOCX → PDF": (["docx"], "pdf"),
    "DOCX → PPTX": (["docx"], "pptx"),
    "DOCX → TXT": (["docx"], "txt"),
    "PPTX → DOCX": (["pptx"], "docx"),
    "PPTX → PDF": (["pptx"], "pdf"),
    "PPTX → TXT": (["pptx"], "txt"),
    "TXT → DOCX": (["txt"], "docx"),
    "TXT → PPTX": (["txt"], "pptx"),
    "TXT → PDF": (["txt"], "pdf"),
    "Image → DOCX": (["png", "jpg", "jpeg", "webp", "tif", "tiff"], "docx"),
    "Image → PDF": (["png", "jpg", "jpeg", "webp", "tif", "tiff"], "pdf"),
    "Image → PPTX": (["png", "jpg", "jpeg", "webp", "tif", "tiff"], "pptx"),
    "Image → TXT": (["png", "jpg", "jpeg", "webp", "tif", "tiff"], "txt"),
}

with st.sidebar:
    st.header("Settings")
    conversion = st.selectbox("Choose conversion", list(CONVERSIONS.keys()))
    ocr_enabled = True  # OCR is on automatically; no checkbox required.
    st.caption("OCR is enabled automatically for scanned PDFs and images.")

    use_ai = st.checkbox(
        "Use AI to correct OCR/text errors (may change text)",
        value=False,
        help="For maximum fidelity, leave this off. AI may alter spelling, numbers, or punctuation.",
    )
    groq_key = st.text_input(
        "Groq API key",
        type="password",
        help="Only required when AI cleanup is enabled.",
    )
    model_name = st.selectbox(
        "Groq model",
        ["llama-3.3-70b-versatile"],
        disabled=not use_ai,
    )

accepted_extensions, output_format = CONVERSIONS[conversion]
uploaded = st.file_uploader(
    f"Upload a file for {conversion}",
    type=accepted_extensions,
)
st.info(f"Selected conversion: **{conversion}** · Output format: **.{output_format}**")


# Few-shot examples are deliberately conservative. AI cleanup is opt-in because it
# cannot guarantee verbatim preservation.
EXAMPLES = [
    {
        "input_text": "OCR text: The quick br0wn fox jumps over the l0zy dog.",
        "output_text": "The quick brown fox jumps over the lazy dog.",
    },
    {
        "input_text": "OCR text: COMPUTER NETWORKS\nA network connects devices and allow data shanng.",
        "output_text": "COMPUTER NETWORKS\n\nA network connects devices and allows data sharing.",
    },
    {
        "input_text": "OCR text: Renewable energy comes from naturaly replenished sources.",
        "output_text": "Renewable energy comes from naturally replenished sources.",
    },
]

if LANGCHAIN_AVAILABLE:
    example_prompt = PromptTemplate(
        input_variables=["input_text", "output_text"],
        template="Input example:\n{input_text}\n\nExpected cleaned output:\n{output_text}",
    )
    few_shot_prompt = FewShotPromptTemplate(
        examples=EXAMPLES,
        example_prompt=example_prompt,
        prefix=(
            "Correct only obvious OCR mistakes. Preserve all names, numbers, dates, "
            "formulas, headings, lists, punctuation, and paragraph breaks unless an "
            "error is unmistakable. Do not summarize or add information. Return only text."
        ),
        suffix="Clean this extracted text:\n\n{document_text}\n\nCorrected text:",
        input_variables=["document_text"],
    )
else:
    few_shot_prompt = None


def run_ocr_on_image(image):
    return pytesseract.image_to_string(image, lang="eng")


def extract_pdf_text(file_bytes, use_ocr=True):
    pdf = fitz.open(stream=file_bytes, filetype="pdf")
    pages = []
    for page_number, page in enumerate(pdf, start=1):
        text = page.get_text("text").strip()
        if not text and use_ocr:
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image = Image.open(io.BytesIO(pix.tobytes("png")))
            text = run_ocr_on_image(image).strip()
        pages.append(text)
    pdf.close()
    return "\n\n".join(pages)


def extract_docx_text(file_bytes):
    doc = Document(io.BytesIO(file_bytes))
    chunks = []
    for paragraph in doc.paragraphs:
        chunks.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            chunks.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(chunks)


def extract_pptx_text(file_bytes):
    prs = Presentation(io.BytesIO(file_bytes))
    chunks = []
    for slide_number, slide in enumerate(prs.slides, start=1):
        chunks.append(f"[Slide {slide_number}]")
        # Reading order by position is more useful than arbitrary XML order.
        shapes = sorted(slide.shapes, key=lambda s: (getattr(s, "top", 0), getattr(s, "left", 0)))
        for shape in shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                chunks.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    chunks.append(" | ".join(cell.text for cell in row.cells))
    return "\n\n".join(chunks)


def extract_text_file(file_bytes):
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="replace")


def extract_image_text(file_bytes):
    image = Image.open(io.BytesIO(file_bytes))
    return run_ocr_on_image(image)


def clean_with_groq(text, api_key, model):
    if not LANGCHAIN_AVAILABLE:
        raise RuntimeError(
            "AI cleanup dependencies are missing. Install langchain-core and langchain-groq."
        )
    api_key = api_key or os.environ.get("GROQ_API_KEY", "")
    if not api_key:
        raise ValueError("Enter a Groq API key or configure GROQ_API_KEY in Streamlit Secrets.")
    llm = ChatGroq(api_key=api_key, model=model, temperature=0)
    chain = few_shot_prompt | llm
    chunks = [text[i:i + 10000] for i in range(0, len(text), 10000)] or [""]
    return "\n\n".join(chain.invoke({"document_text": chunk}).content.strip() for chunk in chunks)


def split_sections(text):
    blocks = [b.strip() for b in re.split(r"\n\s*\n", text) if b.strip()]
    return blocks or [text.strip()]


def make_docx_from_text(text):
    """Used for TXT/OCR output only; it cannot reproduce an original layout."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)

    for line in text.splitlines():
        if not line.strip():
            doc.add_paragraph("")
        elif re.match(r"^\s*[-•*]\s+", line):
            doc.add_paragraph(re.sub(r"^\s*[-•*]\s+", "", line), style="List Bullet")
        else:
            doc.add_paragraph(line)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def make_pptx_from_text(text):
    """Text-to-slide conversion is a layout reconstruction, not a pixel-perfect conversion."""
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        lines = ["No readable text was extracted."]
    for start in range(0, len(lines), 5):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Converted Document" if start == 0 else f"Content {start // 5 + 1}"
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for index, line in enumerate(lines[start:start + 5]):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.level = 0
            paragraph.font.size = PptPt(20)
    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()


def make_pdf_from_text(text):
    output = io.BytesIO()
    pdf = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=0.65 * inch,
        leftMargin=0.65 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.65 * inch,
    )
    styles = getSampleStyleSheet()
    story = []
    for line in text.splitlines():
        if not line.strip():
            story.append(Spacer(1, 7))
            continue
        safe = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph(safe, styles["BodyText"]))
        story.append(Spacer(1, 4))
    pdf.build(story)
    return output.getvalue()


def find_libreoffice():
    return shutil.which("libreoffice") or shutil.which("soffice")


def libreoffice_convert(file_bytes, source_ext, target_ext):
    """Use LibreOffice's real document renderer for DOCX/PPTX/PDF conversions."""
    executable = find_libreoffice()
    if not executable:
        raise RuntimeError(
            f"LibreOffice is required for {source_ext.upper()} → {target_ext.upper()} "
            "to preserve layout. Install LibreOffice on the server, or use a deployment "
            "environment that includes it."
        )
    with tempfile.TemporaryDirectory() as temp_dir:
        input_path = Path(temp_dir) / f"input.{source_ext}"
        input_path.write_bytes(file_bytes)
        profile_path = Path(temp_dir) / "lo-profile"
        command = [
            executable,
            f"-env:UserInstallation={profile_path.as_uri()}",
            "--headless",
            "--convert-to",
            target_ext,
            "--outdir",
            temp_dir,
            str(input_path),
        ]
        result = subprocess.run(command, capture_output=True, text=True, timeout=180)
        output_path = Path(temp_dir) / f"input.{target_ext}"
        if result.returncode != 0 or not output_path.exists():
            details = (result.stderr or result.stdout or "No details returned").strip()
            raise RuntimeError(f"LibreOffice conversion failed: {details}")
        return output_path.read_bytes()


def pdf_to_docx(file_bytes):
    """Use pdf2docx for selectable-text PDFs; OCR fallback for scanned PDFs."""
    if not PDF2DOCX_AVAILABLE:
        raise RuntimeError("Install pdf2docx to convert PDF to editable DOCX.")
    with tempfile.TemporaryDirectory() as temp_dir:
        source = Path(temp_dir) / "source.pdf"
        target = Path(temp_dir) / "converted.docx"
        source.write_bytes(file_bytes)
        converter = PDFConverter(str(source))
        try:
            converter.convert(str(target), start=0, end=None)
        finally:
            converter.close()

        # Detect whether the PDF has meaningful selectable text.
        pdf = fitz.open(stream=file_bytes, filetype="pdf")
        selectable_text = sum(len(page.get_text("text").strip()) for page in pdf)
        pdf.close()

        if selectable_text > 20 and target.exists():
            return target.read_bytes()

    # Scanned PDF fallback: OCR then create an editable text document.
    ocr_text = extract_pdf_text(file_bytes, use_ocr=True)
    if not ocr_text.strip():
        raise RuntimeError("No readable text found. Check the PDF quality and OCR installation.")
    return make_docx_from_text(ocr_text)


def pdf_to_pptx_as_page_images(file_bytes):
    """Preserves the appearance of every PDF page; page content is image-based, not editable text."""
    pdf = fitz.open(stream=file_bytes, filetype="pdf")
    prs = Presentation()
    prs.slide_width = PptInches(10)
    prs.slide_height = PptInches(7.5)
    for page in pdf:
        pix = page.get_pixmap(matrix=fitz.Matrix(1.8, 1.8), alpha=False)
        image_bytes = pix.tobytes("png")
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        page_width = page.rect.width
        page_height = page.rect.height
        slide_ratio = prs.slide_width / prs.slide_height
        page_ratio = page_width / page_height
        if page_ratio > slide_ratio:
            width = prs.slide_width
            height = int(width / page_ratio)
            left, top = 0, int((prs.slide_height - height) / 2)
        else:
            height = prs.slide_height
            width = int(height * page_ratio)
            left, top = int((prs.slide_width - width) / 2), 0
        slide.shapes.add_picture(io.BytesIO(image_bytes), left, top, width=width, height=height)
    pdf.close()
    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()


def image_to_pdf(file_bytes):
    image = Image.open(io.BytesIO(file_bytes)).convert("RGB")
    output = io.BytesIO()
    image.save(output, format="PDF", resolution=150.0)
    return output.getvalue()


def image_to_pptx(file_bytes):
    image = Image.open(io.BytesIO(file_bytes)).convert("RGB")
    prs = Presentation()
    prs.slide_width = PptInches(10)
    prs.slide_height = PptInches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ratio = image.width / image.height
    slide_ratio = prs.slide_width / prs.slide_height
    if ratio > slide_ratio:
        width = prs.slide_width
        height = int(width / ratio)
        left, top = 0, int((prs.slide_height - height) / 2)
    else:
        height = prs.slide_height
        width = int(height * ratio)
        left, top = int((prs.slide_width - width) / 2), 0
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    slide.shapes.add_picture(io.BytesIO(buffer.getvalue()), left, top, width=width, height=height)
    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()


def pptx_to_docx(file_bytes):
    """Retains text, basic paragraphs and tables; original slide positioning is not representable in DOCX."""
    prs = Presentation(io.BytesIO(file_bytes))
    doc = Document()
    for slide_number, slide in enumerate(prs.slides, start=1):
        doc.add_heading(f"Slide {slide_number}", level=1)
        shapes = sorted(slide.shapes, key=lambda s: (getattr(s, "top", 0), getattr(s, "left", 0)))
        for shape in shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                for paragraph in shape.text_frame.paragraphs:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        r = p.add_run(run.text)
                        if run.font.bold is not None:
                            r.bold = run.font.bold
                        if run.font.italic is not None:
                            r.italic = run.font.italic
                        if run.font.size is not None:
                            r.font.size = Pt(run.font.size.pt)
                        if run.font.name:
                            r.font.name = run.font.name
            if getattr(shape, "has_table", False):
                table = doc.add_table(rows=0, cols=len(shape.table.columns))
                table.style = "Table Grid"
                for source_row in shape.table.rows:
                    cells = table.add_row().cells
                    for i, source_cell in enumerate(source_row.cells):
                        cells[i].text = source_cell.text
            if shape.shape_type == 13:  # picture
                try:
                    image_stream = io.BytesIO(shape.image.blob)
                    doc.add_picture(image_stream, width=Inches(5.8))
                except Exception:
                    pass
        if slide_number < len(prs.slides):
            doc.add_page_break()
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def convert_file(conversion_name, file_bytes, source_ext, target_ext):
    if conversion_name == "PDF → DOCX":
        return pdf_to_docx(file_bytes), (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "pdf2docx reconstructs editable content; complex layouts may still differ.",
        )
    if conversion_name == "PDF → PPTX":
        return pdf_to_pptx_as_page_images(file_bytes), (
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "Each PDF page is a slide image for high visual fidelity; its text is not directly editable.",
        )
    if conversion_name == "PDF → TXT":
        return extract_pdf_text(file_bytes, use_ocr=True).encode("utf-8"), (
            "text/plain; charset=utf-8", "Text output cannot preserve page layout or styling."
        )
    if conversion_name == "DOCX → PDF":
        return libreoffice_convert(file_bytes, "docx", "pdf"), (
            "application/pdf", "Rendered with LibreOffice to preserve Word layout as closely as possible."
        )
    if conversion_name == "PPTX → PDF":
        return libreoffice_convert(file_bytes, "pptx", "pdf"), (
            "application/pdf", "Rendered with LibreOffice; results can vary if fonts are missing."
        )
    if conversion_name == "DOCX → PPTX":
        text = extract_docx_text(file_bytes)
        return make_pptx_from_text(text), (
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "Text and section order are retained, but Word page layout cannot map perfectly to slides."
        )
    if conversion_name == "DOCX → TXT":
        return extract_docx_text(file_bytes).encode("utf-8"), (
            "text/plain; charset=utf-8", "Plain text cannot preserve Word formatting."
        )
    if conversion_name == "PPTX → DOCX":
        return pptx_to_docx(file_bytes), (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "Text, basic formatting, tables and pictures are transferred; exact slide positioning is not retained."
        )
    if conversion_name == "PPTX → TXT":
        return extract_pptx_text(file_bytes).encode("utf-8"), (
            "text/plain; charset=utf-8", "Plain text cannot preserve slide layout or styling."
        )
    if conversion_name == "TXT → DOCX":
        return make_docx_from_text(extract_text_file(file_bytes)), (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "Text and line breaks are retained; a TXT file has no font/layout metadata."
        )
    if conversion_name == "TXT → PPTX":
        return make_pptx_from_text(extract_text_file(file_bytes)), (
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "Text is divided into readable slides; a TXT file has no slide-layout metadata."
        )
    if conversion_name == "TXT → PDF":
        return make_pdf_from_text(extract_text_file(file_bytes)), (
            "application/pdf", "Text and line order are retained; a TXT file has no page-layout metadata."
        )
    if conversion_name == "Image → DOCX":
        return make_docx_from_text(extract_image_text(file_bytes)), (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "OCR creates editable text; exact placement and typography require manual review."
        )
    if conversion_name == "Image → PDF":
        return image_to_pdf(file_bytes), (
            "application/pdf", "The image appearance is preserved on the PDF page."
        )
    if conversion_name == "Image → PPTX":
        return image_to_pptx(file_bytes), (
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "The image appearance is preserved on the slide; text within it is not editable."
        )
    if conversion_name == "Image → TXT":
        return extract_image_text(file_bytes).encode("utf-8"), (
            "text/plain; charset=utf-8", "OCR text may contain recognition errors and needs review."
        )
    raise ValueError(f"Unsupported conversion: {conversion_name}")


if uploaded:
    input_ext = Path(uploaded.name).suffix.lower().lstrip(".")
    if st.button("🚀 Convert file", type="primary", use_container_width=True):
        try:
            file_bytes = uploaded.getvalue()
            with st.spinner("Converting with the appropriate document tools..."):
                output, (mime, fidelity_note) = convert_file(
                    conversion, file_bytes, input_ext, output_format
                )

            # AI cleanup is optional and only applied to text-centric conversions.
            if use_ai:
                if conversion not in {"PDF → TXT", "DOCX → TXT", "PPTX → TXT", "Image → TXT"}:
                    st.warning(
                        "AI cleanup is skipped for this format to avoid altering document layout/content."
                    )
                else:
                    extracted_text = output.decode("utf-8", errors="replace")
                    cleaned_text = clean_with_groq(extracted_text, groq_key, model_name)
                    output = cleaned_text.encode("utf-8")

            st.success("Conversion completed.")
            st.info(fidelity_note)
            stem = Path(uploaded.name).stem
            output_name = f"{stem}_converted.{output_format}"
            st.download_button(
                label=f"⬇️ Download {output_name}",
                data=output,
                file_name=output_name,
                mime=mime,
                use_container_width=True,
            )
        except Exception as exc:
            st.error(f"Conversion failed: {exc}")
            st.info(
                "For best fidelity, use the original file type when possible. "
                "For DOCX/PPTX to PDF, make sure LibreOffice is installed in the deployment environment."
            )
else:
    st.subheader("How to use")
    st.markdown(
        "1. Choose a conversion from the sidebar.\n"
        "2. Upload a supported file.\n"
        "3. Click **Convert file** and download the result.\n\n"
        "OCR runs automatically for image files and scanned PDF pages. "
        "For exact text preservation, leave AI cleanup disabled."
    )
